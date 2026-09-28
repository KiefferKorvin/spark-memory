import asyncio
import base64
import io
import json
import socket

import httpx
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from graph_memory.api import create_app
from graph_memory.config import Settings
from graph_memory.demo import DemoExternal, DemoModels, DOCUMENTS, QUESTIONS, seed
from graph_memory.evidence import ContextBuilder, SufficiencyEvaluator, model_view
from graph_memory.graph import InMemoryGraph
from graph_memory.llm import OpenRouter, ProviderError
from graph_memory.models import (Concept, Coverage, CoverageItem, Edge, Evidence, IngestRequest,
                                 NavigationDecision, Relation, Understanding)
from graph_memory.parsing import chunk_text, extract, parse_structure, token_count
from graph_memory.service import Memory
from graph_memory.sources import SafeFetcher


def settings(**kwargs):
    return Settings(memory_mode="demo", _env_file=None, **kwargs)


def make_memory(**kwargs):
    return Memory(settings(**kwargs), InMemoryGraph(), DemoModels(), DemoExternal())


def test_hierarchy_and_small_document_leaf():
    parsed = parse_structure("Brief source.", "note", "small", settings())
    assert parsed.document.retrieval_leaf and len(parsed.nodes) == 1
    text = "# A\nintro\n## B\nb\n1.2.3.4.5.6.7 Deep\ndeep text\n## C\nlast"
    parsed = parse_structure(text, "Book", "doc", settings())
    sections = [n for n in parsed.nodes if n.kind == "Section"]
    deep = next(n for n in sections if n.level == 7)
    assert deep.section_path == ["A", "B", "1.2.3.4.5.6.7 Deep"]
    assert next(n for n in sections if n.label == "C").section_path == ["A", "C"]
    assert any(n.context_header == "Book > A > B > 1.2.3.4.5.6.7 Deep" for n in parsed.nodes if n.kind == "Chunk")
    code = parse_structure("```python\n# not a heading\n1.2.3 not either\n```", "code", "code", settings())
    assert not any(n.kind == "Section" for n in code.nodes)


@pytest.mark.parametrize("text", ["One sentence. "*300, "你好世界🌍"*200, "X"*10000, "Paragraph one.\n\n"*150], ids=["sentences", "unicode", "long_word", "paragraphs"])
def test_chunk_boundaries_unicode_and_token_limit(text):
    chunks = chunk_text(text, 35, 50)
    assert "".join(chunks) == text
    assert all(token_count(c) <= 50 for c in chunks)


def test_html_json_email_and_docx_parsing():
    assert "bad()" not in extract(b"<h1>Title</h1><p>Body</p><script>bad()</script>", "text/html", 10000)
    assert json.loads(extract(b'{"x":1}', "application/json", 1000)) == {"x": 1}
    email = b"Subject: Test\nFrom: a@example.org\nContent-Type: text/plain; charset=utf-8\n\nA message"
    assert "A message" in extract(email, "message/rfc822", 10000)
    from docx import Document
    doc = Document(); doc.add_heading("Outer", 1); doc.add_heading("Inner", 2); doc.add_paragraph("Original text.")
    output = io.BytesIO(); doc.save(output)
    text = extract(output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 100000)
    assert "# Outer" in text and "## Inner" in text and "Original text." in text


async def test_concept_reuse_duplicate_ingestion_multiple_provenance():
    memory = make_memory()
    first = await memory.ingest(IngestRequest(**DOCUMENTS[0]))
    count = len(memory.repository.nodes)
    again = await memory.ingest(IngestRequest(**DOCUMENTS[0]))
    assert again["duplicate"] and again["document_id"] == first["document_id"]
    assert len(memory.repository.nodes) == count
    await memory.ingest(IngestRequest(title="Second source", text=DOCUMENTS[0]["text"], mime_type="text/markdown"))
    assert len((await memory.repository.provenance(first["document_id"]))["sources"]) == 2
    await memory.ingest(IngestRequest(title="Piano", text="Piano and harmony work together."))
    assert len([n for n in memory.repository.nodes.values() if n.kind == "Concept" and n.label == "Piano"]) == 1
    rootless = await memory.repository.exact_concept("Rootless voicings")
    parents = [e for e in memory.repository.edges.values() if e.target == rootless.id and e.relation == Relation.BROADER_THAN]
    assert len(parents) == 2


async def test_alias_resolution_cycle_rollback_and_external_alignment():
    memory = make_memory()
    a = Concept(id="a", label="Keyboard", preferred_label="Keyboard", aliases=["Piano"])
    b = Concept(id="b", label="Instrument", preferred_label="Instrument")
    await memory.repository.put([a, b], [Edge(source="b", target="a", relation=Relation.BROADER_THAN)])
    assert (await memory.repository.exact_concept("piano")).id == "a"
    with pytest.raises(ValueError, match="cycle"):
        await memory.repository.put([Concept(id="c", label="new", preferred_label="new")], [Edge(source="a", target="b", relation=Relation.BROADER_THAN)])
    assert await memory.repository.get("c") is None
    await memory.ingestion.ontology.align("a", "demo", "123", "External keyboard", exact=True)
    assert any(e.relation == Relation.EXACT_MATCH for e in memory.repository.edges.values())


async def test_hybrid_candidate_generation_ranks_before_limit():
    memory = make_memory(candidate_limit=3)
    root = Concept(id="root", label="Root", preferred_label="Root")
    nodes = [Concept(id=f"n{i}", label="unrelated" if i < 100 else "rootless piano", preferred_label=str(i)) for i in range(101)]
    await memory.repository.put([root, *nodes], [Edge(source="root", target=n.id, relation=Relation.BROADER_THAN) for n in nodes])
    result = await memory.repository.candidates("rootless piano", await memory.models.embed("rootless piano"), 3, parent="root")
    assert len(result) == 3 and result[0][0].id == "n100"


async def test_end_to_end_parallel_loop_prevention_pruning_external_ingestion_and_trace():
    memory = make_memory()
    await memory.initialize()
    await seed(memory)
    result = await memory.query(QUESTIONS[3])
    assert result["status"] == "completed", result
    assert result["coverage"]["overall_status"] == "SUFFICIENT", result["coverage"]
    events = await memory.events(result["query_id"])
    types = {e["event_type"] for e in events}
    assert {"NODE_PRUNED", "BRANCH_SPAWNED", "EVIDENCE_FOUND", "EXTERNAL_SEARCH_STARTED", "SOURCE_INGESTED", "QUERY_COMPLETED"} <= types
    assert len({e["event_id"] for e in events}) == len(events)
    assert [e["sequence"] for e in events] == list(range(1, len(events)+1))
    explored = [(e["information_need_id"], e["node_id"]) for e in events if e["event_type"] in {"NODE_SELECTED", "NODE_EXPANDED"}]
    assert len(explored) == len(set(explored))
    assert all(len(e["metadata"]["nodes"]) <= memory.settings.candidate_limit for e in events if e["event_type"] == "CANDIDATES_GENERATED")
    assert any(s["metadata"].get("original_uri") for e in result["evidence"] for s in e["provenance"]["sources"])
    assert result["nodes_explored"] <= memory.settings.max_total_nodes_explored
    await memory.close()


async def test_budget_and_missing_coverage_without_external():
    from graph_memory.models import QueryRequest
    memory = make_memory(max_total_nodes_explored=2, max_depth=1)
    await seed(memory)
    result = await memory.query(QueryRequest(query="unknown orbital period", allow_external=False))
    assert result["nodes_explored"] <= 2
    # Coverage only decides whether to search online, so it is not assessed (or paid for) when that is disabled.
    assert result["coverage"] is None
    events = await memory.events(result["query_id"])
    assert not any(e["event_type"] in ("EXTERNAL_SEARCH_STARTED", "SUFFICIENCY_CHECKED") for e in events)
    assert any(e["event_type"] == "EXTERNAL_SEARCH_SKIPPED" for e in events)


async def test_enrichment_preserves_support_and_is_idempotent():
    memory = make_memory()
    result = await memory.ingest(IngestRequest(title="Piano", text="Piano has keys."))
    nodes = await memory.enrichment.enrich(result["document_id"], force=True)
    assert len(nodes) == 1 and nodes[0].assertion_type == "claim"
    assert (await memory.repository.provenance(nodes[0].id))["sources"]
    assert await memory.enrichment.enrich(result["document_id"], force=True) == []


async def test_coverage_rejects_hallucinated_evidence_and_context_budget():
    from graph_memory.models import Need
    needs = [Need(id="N1", description="topic")]
    item = Evidence(id="E1", information_need_ids=["N1"], source_node_id="a", source_type="chunk", text="original "*2000,
                    relevance_score=.9, confidence=.8, provenance={"sources": []})
    context = ContextBuilder(300).build(needs, [item, item])
    assert token_count(json.dumps(model_view(context)[0], ensure_ascii=False)) <= 300  # the budget counts what models read
    assert len(context) == 1 and context[0]["text"] in item.text
    view = [{"id": "E1", "text": item.text}]
    class BadCoverage(DemoModels):
        async def structured(self, operation, payload, schema, query_id=None):
            return Coverage(coverage=[CoverageItem(information_need_id="N1", status="COVERED", evidence_ids=["invented"])], overall_status="SUFFICIENT")
    # An invented ID supports nothing, and a need the model skipped stays MISSING, without discarding valid verdicts.
    bad = await SufficiencyEvaluator(BadCoverage()).evaluate([*needs, Need(id="N2", description="other")], view, "q")
    assert [(c.status, c.evidence_ids) for c in bad.coverage] == [("MISSING", []), ("MISSING", [])]
    assert bad.overall_status == "INSUFFICIENT"
    class CrossNeed(DemoModels):
        async def structured(self, operation, payload, schema, query_id=None):
            return Coverage(coverage=[CoverageItem(information_need_id=n, status="COVERED", evidence_ids=["E1"]) for n in ("N1", "N2")], overall_status="SUFFICIENT")
    # E1 was found for N1 but also answers N2; that is valid support, not a hallucination.
    both = [*needs, Need(id="N2", description="other")]
    assert (await SufficiencyEvaluator(CrossNeed()).evaluate(both, view, "q")).overall_status == "SUFFICIENT"


def test_ssrf_checks_public_dns_pinning_and_input_validation(monkeypatch):
    fetcher = SafeFetcher(["allowed.example"], 1024)
    for url in ["http://allowed.example/", "https://127.0.0.1/", "https://allowed.example:444/", "https://user:pass@allowed.example/"]:
        with pytest.raises(ValueError):
            fetcher.validate(url)
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(2,1,6,"",("127.0.0.1",443))])
    with pytest.raises(ValueError, match="Non-public"):
        fetcher.validate("https://allowed.example")
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: [(2,1,6,"",("8.8.8.8",443))])
    assert fetcher.validate("https://allowed.example")[1] == "8.8.8.8"
    with pytest.raises(ValidationError):
        IngestRequest(title="x", text="a", url="https://example.org")
    with pytest.raises(ValidationError):
        NavigationDecision(decisions=[{"node_id":"x", "action":"DELETE"}], current_node_action="CONTINUE")


async def test_openrouter_jev_protocol_schema_repair_and_usage():
    requests = []
    async def handler(request):
        data = json.loads(request.content); requests.append((request.url.path, data))
        if request.url.path.endswith("decisions"):
            return httpx.Response(200, json={"answers": {key:{"type":"choice", "choice":"CONTINUE" if key == "branch_status" else "SELECT"} for key in data["questions"]}, "usage":{"input_tokens":20,"output_tokens":0,"cost":.001}})
        if len([r for r in requests if r[0].endswith("completions")]) == 1:
            return httpx.Response(200, json={"choices":[{"message":{"content":"not json"}}]})
        return httpx.Response(200, json={"choices":[{"message":{"content":json.dumps({"document_type":"note","summary":"s","routing_summary":"r"})}}]})
    repo = InMemoryGraph()
    provider = OpenRouter(settings(provider_retries=1), repo, httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    result = await provider.navigate({"candidate_children":[{"node_id":"n"}]}, "q")
    assert result.decisions[0].action == "SELECT"
    assert requests[0][0] == "/api/alpha/decisions" and requests[0][1]["model"] == "typesafe/jev-1.13"
    state = requests[0][1]["state"]  # the policy travels once; answers map back by index, so no node IDs
    assert state["navigation_policy"] and state["candidate_children"] == [{}]
    assert all(len(q["instructions"]) < 100 for q in requests[0][1]["questions"].values())
    await provider.structured("understanding", {"text":"x"}, Understanding)
    assert len(await repo.records("usage")) == 3
    await provider.structured("understanding", {"text":"x"}, Understanding)
    assert len(requests) == 3  # prompt-version/model keyed cache
    await provider.close()


async def test_parallel_workers_are_bounded_and_provider_failure_is_visible():
    class ObservedModels(DemoModels):
        active = 0
        peak = 0
        async def navigate(self, payload, query_id):
            self.active += 1
            self.peak = max(self.peak, self.active)
            try:
                await asyncio.sleep(.01)
                return await super().navigate(payload, query_id)
            finally:
                self.active -= 1
    models = ObservedModels()
    memory = Memory(settings(max_parallel_branches=2, max_total_nodes_explored=7), InMemoryGraph(), models)
    await seed(memory)
    result = await memory.query("How do piano and harmony connect?")
    assert result["status"] == "completed" and result["nodes_explored"] <= 7
    assert models.peak == 2
    class FailedNavigation(DemoModels):
        async def navigate(self, payload, query_id):
            raise ProviderError("unavailable")
    memory.models = FailedNavigation()
    memory.retrieval.models = memory.models
    result = await memory.query("piano keys")
    assert result["status"] == "completed"
    assert any(e["event_type"] == "MODEL_FAILURE" for e in await memory.events(result["query_id"]))


async def test_contradictions_coexist_and_invalid_claim_rolls_back():
    from graph_memory.models import Assertion, Enrichment
    memory = make_memory()
    one = await memory.ingest(IngestRequest(title="source one", text="The value is 10."))
    two = await memory.ingest(IngestRequest(title="source two", text="The value is 20."))
    a = Assertion(id="a", label="First claim", text="The value is 10.", proposition="The value is 10.", confidence=.8)
    b = Assertion(id="b", label="Second claim", text="The value is 20.", proposition="The value is 20.", confidence=.7)
    await memory.repository.put([a,b], [Edge(source="a",target=one["document_id"],relation=Relation.SUPPORTED_BY),
                                      Edge(source="b",target=two["document_id"],relation=Relation.SUPPORTED_BY),
                                      Edge(source="a",target="b",relation=Relation.CONTRADICTS)])
    assert (await memory.repository.get("a")).text != (await memory.repository.get("b")).text
    assert (await memory.repository.provenance("b"))["contradictions"] == ["a"]
    class InvalidQuote(DemoModels):
        async def structured(self, operation, payload, schema, query_id=None):
            return Enrichment(assertions=[{"proposition":"invented", "assertion_type":"claim", "confidence":.9,
                                           "supporting_quote":"not in source", "contradicts_ids":[]}])
    memory.enrichment.models = InvalidQuote()
    count = len(memory.repository.nodes)
    with pytest.raises(ValueError, match="quote"):
        await memory.enrichment.enrich(one["document_id"], force=True)
    assert len(memory.repository.nodes) == count


async def test_concurrent_submit_capacity_and_query_timeout():
    class Slow(DemoModels):
        async def structured(self, *args, **kwargs):
            await asyncio.sleep(2)
            return await super().structured(*args, **kwargs)
    memory = Memory(settings(max_active_queries=1, query_timeout_seconds=1), InMemoryGraph(), Slow())
    query_id = await memory.submit("question")
    with pytest.raises(RuntimeError, match="capacity"):
        await memory.submit("other")
    task = memory.tasks[query_id]
    await task
    result = await memory.repository.read_record("query", query_id)
    assert result["status"] == "failed" and "timed out" in result["error"]


async def test_credit_reserved_by_in_flight_requests_is_retried():
    calls = []
    async def handler(request):
        calls.append(request)
        if len(calls) == 1:
            return httpx.Response(402, headers={"Retry-After": "0.01"}, json={"error": {"message": "This request would exceed your available credits given your current in-flight requests. Retry after in-flight requests settle, or add credits."}})
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"document_type": "note", "summary": "s", "routing_summary": "r"})}}]})
    provider = OpenRouter(settings(provider_retries=1), InMemoryGraph(), httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    assert (await provider.structured("understanding", {"text": "x"}, Understanding)).summary == "s" and len(calls) == 2
    await provider.close()


async def test_bad_provider_choices_retry_then_fail():
    calls = []
    async def bad(request):
        calls.append(request)
        return httpx.Response(200, json={"answers":{"node_0":{"type":"choice","choice":"EXECUTE"}}})
    provider = OpenRouter(settings(provider_retries=1), InMemoryGraph(), httpx.AsyncClient(transport=httpx.MockTransport(bad)))
    with pytest.raises(ProviderError):
        await provider.navigate({"candidate_children":[{"node_id":"n"}]}, "q")
    assert len(calls) == 2
    await provider.close()


def test_api_auth_ingestion_query_events_replay_and_limits():
    memory = make_memory(memory_api_token="test-secret", max_source_bytes=1024)
    with TestClient(create_app(memory)) as client:
        assert client.get("/memory/health").status_code == 401
        client.headers["Authorization"] = "Bearer test-secret"
        assert client.post("/memory/ingest", json={"title":"Piano", "text":"Piano keys create notes."}).status_code == 200
        response = client.post("/memory/query", json={"query":"Piano keys", "allow_external":False})
        assert response.status_code == 202
        query_id = response.json()["query_id"]
        with client.stream("GET", f"/memory/query/{query_id}/stream") as response:
            body = response.read().decode()
        assert "QUERY_COMPLETED" in body and "event: done" in body
        events = client.get(f"/memory/query/{query_id}/events").json()
        resumed = client.get(f"/memory/query/{query_id}/events?after=2").json()
        assert resumed == events[2:]
        assert client.get(f"/memory/query/{query_id}").json()["status"] == "completed"
        assert client.post("/memory/ingest", content=b"x"*20000).status_code == 413
        assert client.get("/memory/nodes/missing").status_code == 404
        assert client.post("/memory/ingest", json={"title":"invalid", "content_base64":base64.b64encode(b'bad PDF').decode(), "mime_type":"application/pdf"}).status_code == 422


async def hub_graph():
    """Six retrievable leaves about one concept hub."""
    from graph_memory.models import Document
    repo = InMemoryGraph()
    docs = [Document(id=f"d{i}", label=f"note {i}", text=f"Piano note number {i}.", retrieval_leaf=True) for i in range(6)]
    await repo.put([Concept(id="hub", label="Piano", preferred_label="Piano"), *docs],
                   [Edge(source=d.id, target="hub", relation=Relation.ABOUT) for d in docs])
    return repo


async def test_root_breadth_and_navigation_floor():
    from graph_memory.models import Branch, Decision, Need
    from graph_memory.retrieval import ExplorationSupervisor
    from graph_memory.trace import TraceService

    class Policy(DemoModels):
        def __init__(self, prune_all=False):
            self.prune_all = prune_all
        async def navigate(self, payload, query_id):
            return NavigationDecision(decisions=[Decision(node_id=c["node_id"], action="SELECT" if c["retrievable"] and not self.prune_all else "PRUNE")
                                                 for c in payload["candidate_children"]], current_node_action="DEAD_END" if self.prune_all else "CONTINUE")

    need = Need(id="N1", description="piano note")
    repo = await hub_graph()
    supervisor = ExplorationSupervisor(settings(max_root_children=8, max_children_per_decision=4), repo, Policy(), TraceService(repo, "q"))
    supervisor.ceiling = 100
    root = Branch(information_need_id="N1")
    await supervisor.work(root, need, [])
    assert len([b for b in supervisor.branches if b.parent_branch_id == root.id]) == 6  # all six root SELECTs, cap 8
    supervisor = ExplorationSupervisor(settings(max_root_children=8, max_children_per_decision=4), repo, Policy(), TraceService(repo, "q"))
    supervisor.ceiling = 100
    hub = Branch(information_need_id="N1", node_id="hub", path=["hub"], depth=1)
    await supervisor.work(hub, need, [])
    assert len([b for b in supervisor.branches if b.parent_branch_id == hub.id]) == 4  # below the root the cap stays 4

    repo = await hub_graph()
    supervisor = ExplorationSupervisor(settings(min_root_children=2), repo, Policy(prune_all=True), TraceService(repo, "q"))
    await supervisor.explore([need])
    events = await repo.records("event", "q:")
    floor = [e for e in events if e["event_type"] == "NAVIGATION_FLOOR"]
    assert len(floor) == 1 and len(floor[0]["metadata"]["nodes"]) == 2
    assert supervisor.explored == 2  # the two top-ranked candidates were still explored


async def test_navigation_excerpts_only_for_retrievable_leaves():
    from graph_memory.models import Assertion, Chunk, Document, Need, Section
    from graph_memory.retrieval import ExplorationSupervisor, preview
    from graph_memory.trace import TraceService
    text = "First boil the chicken in cola, then fry it with flour and egg. " * 30
    nodes = [Document(id="leaf", label="d", text=text, retrieval_leaf=True), Document(id="container", label="c", text=text),
             Chunk(id="chunk", label="k", text=text, token_count=1, section_path=[], context_header="", order=0),
             Assertion(id="assertion", label="a", text=text, proposition=text, confidence=.5),
             Section(id="section", label="s", text=text, level=1, order=0, section_path=[]),
             Concept(id="concept", label="Frying", preferred_label="Frying", description=text)]
    views = {n.id: preview(n, excerpt_tokens=80) for n in nodes}
    assert {i for i, v in views.items() if "excerpt" in v} == {"leaf", "chunk", "assertion"}
    assert text.startswith(views["leaf"]["excerpt"]) and 70 <= token_count(views["leaf"]["excerpt"]) <= 80
    assert "excerpt" not in preview(nodes[0])

    class Spy(DemoModels):
        payloads = []
        async def navigate(self, payload, query_id):
            self.payloads.append(payload)
            return await super().navigate(payload, query_id)
    repo = await hub_graph()
    supervisor = ExplorationSupervisor(settings(navigation_excerpt_tokens=5), repo, Spy(), TraceService(repo, "q"))
    await supervisor.explore([Need(id="N1", description="piano note")])
    root = Spy.payloads[0]["candidate_children"]
    assert all(("excerpt" in c) == (c["kind"] == "Document") for c in root) and any(c["kind"] == "Concept" for c in root)
    assert all(token_count(c["excerpt"]) <= 5 for c in root if "excerpt" in c)


async def test_navigation_request_shortens_excerpts_to_fit_budget():
    from graph_memory.parsing import truncate
    sent = []
    async def handler(request):
        body = json.loads(request.content); sent.append(body)
        return httpx.Response(200, json={"answers": {k: {"type": "choice", "choice": "CONTINUE" if k == "branch_status" else "PRUNE"}
                                                    for k in body["questions"]}})
    def provider(budget):
        return OpenRouter(settings(model_input_token_budget=budget), InMemoryGraph(), httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    # Worst case at candidate_limit=20: every candidate carries a full 80-token excerpt.
    excerpt = truncate(" ".join(f"word{i}" for i in range(400)), 80)
    state = {"information_need": {"id": "N1", "description": "crispy breading"}, "candidate_children": [
        {"node_id": f"n{i}", "label": "note", "routing_summary": "Explore for frying notes."} for i in range(20)]}
    await provider(100000).navigate(state, "q")
    base = token_count(json.dumps(sent[-1], ensure_ascii=False))  # the same request without excerpts
    full = {**state, "candidate_children": [{**c, "excerpt": excerpt} for c in state["candidate_children"]]}
    budget = base + 600  # 20 x 80 excerpt tokens cannot fit; 20 x 20 can
    await provider(budget).navigate(full, "q")
    shortened = sent[-1]
    assert token_count(json.dumps(shortened, ensure_ascii=False)) <= budget
    assert all(0 < token_count(c["excerpt"]) < 80 and excerpt.startswith(c["excerpt"]) for c in shortened["state"]["candidate_children"])


async def test_original_question_reaches_candidate_search_and_relevance():
    from graph_memory.models import QueryRequest
    class Repo(InMemoryGraph):
        queries = []
        async def candidates(self, query, vector, limit, parent=None, kind=None, **scopes):
            self.queries.append((query, bool(vector)))
            return await super().candidates(query, vector, limit, parent, kind, **scopes)
    class Models(DemoModels):
        relevance = []
        async def structured(self, operation, payload, schema, query_id=None):
            if operation == "relevance":
                self.relevance.append(payload)
            return await super().structured(operation, payload, schema, query_id)
    memory = Memory(settings(), Repo(), Models(), DemoExternal())
    await seed(memory)
    Repo.queries.clear()
    question = "How do piano and harmony connect?"  # the demo decomposition rewrites it into two needs
    result = await memory.query(QueryRequest(query=question, allow_external=False))
    assert question not in [n["description"] for n in result["information_needs"]]
    assert (question, True) in Repo.queries  # searched with its own embedding as well as each need
    assert Models.relevance and all(p["question"] == question for p in Models.relevance)


async def test_answer_length_cap_per_query_overrides_the_setting():
    from graph_memory.models import QueryRequest
    class Models(DemoModels):
        limits = []
        async def structured(self, operation, payload, schema, query_id=None):
            if operation == "synthesis":
                self.limits.append(payload.get("answer_max_words"))
            return await super().structured(operation, payload, schema, query_id)
    memory = Memory(settings(answer_max_words=60), InMemoryGraph(), Models())
    await seed(memory)
    await memory.query(QueryRequest(query="rootless voicings", allow_external=False))
    await memory.query(QueryRequest(query="rootless voicings", allow_external=False, answer_max_words=84))
    assert Models.limits == [60, 84]


async def test_overlong_answer_is_rewritten_within_its_limit():
    from graph_memory.models import Answer, QueryRequest
    class Wordy(DemoModels):
        drafts = []
        async def structured(self, operation, payload, schema, query_id=None):
            if operation != "synthesis":
                return await super().structured(operation, payload, schema, query_id)
            evidence_id = payload["evidence"][0]["id"]
            if "draft_answer" in payload:
                self.drafts.append(payload["draft_answer"])
                return Answer(answer=f"Short answer [{evidence_id}].", evidence_ids=[evidence_id])
            return Answer(answer=" ".join(["word"] * 40) + f" [{evidence_id}]", evidence_ids=[evidence_id])
    memory = Memory(settings(), InMemoryGraph(), Wordy())
    await seed(memory)
    result = await memory.query(QueryRequest(query="What is the purpose of rootless voicings?", allow_external=False, answer_max_words=10))
    assert result["answer"].startswith("Short answer") and len(Wordy.drafts) == 1
    event = next(e for e in await memory.events(result["query_id"]) if e["event_type"] == "ANSWER_SHORTENED")
    assert (event["metadata"]["words_before"], event["metadata"]["words_after"]) == (40, 2)  # citations not counted
    Wordy.drafts.clear()
    await memory.query(QueryRequest(query="What is the purpose of rootless voicings?", allow_external=False, answer_max_words=40))
    assert not Wordy.drafts  # within the limit: no rewrite


def test_context_puts_the_most_relevant_evidence_first():
    from graph_memory.models import Need
    def item(i, score):
        return Evidence(id=i, information_need_ids=["N1"], source_node_id=i, source_type="chunk", text=i,
                        relevance_score=score, confidence=1, provenance={"sources": []})
    context = ContextBuilder(5000).build([Need(id="N1", description="x")], [item("weak", .3), item("strong", .9), item("mid", .6)])
    assert [e["id"] for e in context] == ["strong", "mid", "weak"]


async def test_reembedding_after_a_dimension_change_is_resumable():
    from graph_memory.ingestion import embedding_text
    from graph_memory.migrate import reembed
    memory = make_memory(embedding_dimensions=64)  # the demo embedder returns 64 dimensions
    await seed(memory)
    repo = memory.repository
    embedded = [n for n in repo.nodes.values() if n.embedding]
    for node in embedded:  # as if written by an 8-dimension model
        repo.nodes[node.id] = node.model_copy(update={"embedding": [1.0] * 8})
    result = await reembed(repo, memory.models, memory.settings, batch=2)
    assert result["nodes"] == len(embedded) and embedded
    for node in embedded:
        assert repo.nodes[node.id].embedding == await memory.models.embed(embedding_text(repo.nodes[node.id]))
    assert (await reembed(repo, memory.models, memory.settings))["nodes"] == 0


async def test_external_sources_are_explored_with_fresh_budget():
    from graph_memory.models import QueryRequest
    # The first pass exhausts its node budget; newly ingested sources are still examined directly.
    memory = make_memory(max_total_nodes_explored=2)
    await seed(memory)
    result = await memory.query(QUESTIONS[3])
    practice = next(c for c in result["coverage"]["coverage"] if c["information_need_id"] == "N2")
    assert practice["status"] == "COVERED", result["coverage"]
    assert result["nodes_explored"] > 2  # the external pass ran on its own budget
    ingested = next(e["metadata"] for e in await memory.events(result["query_id"]) if e["event_type"] == "SOURCE_INGESTED")
    # Query-time ingestion only parses and embeds; the source is structured once a later query uses it.
    document_id = ingested["document_id"]
    assert ingested["light"] and document_id in {d for e in result["context"] for d in e["provenance"]["documents"]}
    def structured():
        return not memory.repository.nodes[document_id].metadata["light"] and any(
            e.relation == Relation.ABOUT and e.source == document_id for e in memory.repository.edges.values())
    await asyncio.gather(*memory.ingestion.background)
    assert not structured()  # the query that fetched it does not pay for it
    again = await memory.query(QueryRequest(query=QUESTIONS[3], allow_external=False))
    assert document_id in {d for e in again["context"] for d in e["provenance"]["documents"]}
    await asyncio.gather(*memory.ingestion.background)
    assert structured() and memory.repository.nodes[document_id].summary
    sources = (await memory.repository.provenance(document_id))["sources"]
    assert sources and all(s["metadata"]["metadata_method"] == "structured_model" for s in sources)  # dates and authors extracted


async def test_synthesis_reads_short_ids_and_the_answer_cites_real_ones():
    from graph_memory.models import Answer, QueryRequest
    class Citing(DemoModels):
        payloads = []
        async def structured(self, operation, payload, schema, query_id=None):
            if operation != "synthesis":
                return await super().structured(operation, payload, schema, query_id)
            self.payloads.append(payload)
            ids = [e["id"] for e in payload["evidence"]]
            return Answer(answer=f"Both [{', '.join(ids)}].", evidence_ids=ids)
    memory = Memory(settings(), InMemoryGraph(), Citing())
    await seed(memory)
    question = QueryRequest(query="How do piano and harmony connect?", allow_external=False)
    result = await memory.query(question)
    view = Citing.payloads[0]["evidence"]
    assert len(view) >= 2 and [e["id"] for e in view] == [f"E{i}" for i in range(1, len(view) + 1)]
    assert all(set(e) <= {"id", "text", "source", "contradicts"} for e in view)  # no provenance hashes or UUIDs
    assert all(e["source"]["retrieved_at"] for e in view)  # freshness: when the memory fetched each source
    real = [e["id"] for e in result["context"]]
    assert result["answer"] == "Both " + "".join(f"[{i}]" for i in real) + "." and result["evidence_ids"] == real
    Citing.payloads.clear()
    result = await memory.query(question.model_copy(update={"synthesize": False}))
    assert not Citing.payloads and result["answer"] == "" and result["context"]


async def test_expansion_waits_until_selected_leaves_leave_the_need_without_evidence():
    from graph_memory.models import Decision, Need, Relevance
    from graph_memory.retrieval import ExplorationSupervisor
    from graph_memory.trace import TraceService
    class Policy(DemoModels):  # selects every leaf, expands every concept
        async def navigate(self, payload, query_id):
            return NavigationDecision(decisions=[Decision(node_id=c["node_id"], action="SELECT" if c["retrievable"] else "EXPAND")
                                                 for c in payload["candidate_children"]], current_node_action="CONTINUE")
    class Rejecting(Policy):  # ... and finds no leaf relevant
        async def structured(self, operation, payload, schema, query_id=None):
            return Relevance(relevant=False, confidence=.9) if operation == "relevance" else await super().structured(operation, payload, schema, query_id)
    from graph_memory.models import Source
    for models, expanded in ((Policy(), False), (Rejecting(), True)):
        repo = await hub_graph()
        await repo.put([Source(id=f"s{i}", label="s", source_type="api", mime_type="text/plain", content_hash=str(i)) for i in range(6)],
                       [Edge(source=f"s{i}", target=f"d{i}", relation=Relation.PROVIDES) for i in range(6)])
        await ExplorationSupervisor(settings(), repo, models, TraceService(repo, "q")).explore([Need(id="N1", description="piano note")])
        events = await repo.records("event", "q:")
        assert any(e["event_type"] == "NODE_EXPANDED" and e["node_id"] == "hub" for e in events) == expanded
        assert any(e["metadata"].get("reason") == "need_has_evidence" for e in events) != expanded
        # A selected leaf is retrieved, relevant or not, never navigated from.
        assert not any(e["event_type"] == "CANDIDATES_GENERATED" and e["node_id"] not in (None, "hub") for e in events)


async def test_query_embeddings_take_one_request_per_step():
    from graph_memory.models import QueryRequest
    class Counting(DemoModels):
        requests = []
        async def embed(self, text, query_id=None):
            self.requests.append([text])
            return await DemoModels.embed(self, text)
        async def embed_batch(self, texts, query_id=None):
            self.requests.append(texts)
            return [await DemoModels.embed(self, t) for t in texts]
    memory = Memory(settings(), InMemoryGraph(), Counting())
    await seed(memory)
    Counting.requests.clear()
    await memory.query(QueryRequest(query="rootless voicings", allow_external=False))
    assert Counting.requests == [["rootless voicings"]]  # a lone need shares the question's vector
    Counting.requests.clear()
    await memory.query(QueryRequest(query="How do piano and harmony connect?", allow_external=False))
    assert [len(r) for r in Counting.requests] == [1, 2]  # the question beside decomposition, then both needs at once


async def test_failed_decomposition_or_embedding_degrades_instead_of_failing_the_query():
    from graph_memory.models import QueryRequest
    class Flaky(DemoModels):
        async def structured(self, operation, payload, schema, query_id=None):
            if operation == "decomposition":
                raise ProviderError("decomposition failed after 3 attempt(s): ConnectError: name not known")
            return await super().structured(operation, payload, schema, query_id)
        async def embed(self, text, query_id=None):
            raise ProviderError("embedding failed after 3 attempt(s): HTTP 503")
        async def embed_batch(self, texts, query_id=None):
            raise ProviderError("embedding failed after 3 attempt(s): HTTP 503")
    memory = Memory(settings(), InMemoryGraph(), DemoModels())
    await seed(memory)
    memory.retrieval.models = memory.models = Flaky()
    question = "What is the purpose of rootless voicings?"
    result = await memory.query(QueryRequest(query=question, allow_external=False))
    assert result["status"] == "completed" and result["evidence"]  # lexical search still finds it
    assert [n["description"] for n in result["information_needs"]] == [question]
    failures = {(e["metadata"]["operation"], e["metadata"]["fallback"]) for e in await memory.events(result["query_id"])
                if e["event_type"] == "MODEL_FAILURE"}
    assert {("decomposition", "question_as_need"), ("embedding", "lexical_search")} <= failures


async def test_refused_credits_fail_the_query_with_their_reason():
    from graph_memory.llm import AccountError
    async def handler(request):
        return httpx.Response(402, json={"error": {"message": "Insufficient credits"}})
    provider = OpenRouter(settings(provider_retries=2), InMemoryGraph(), httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    with pytest.raises(AccountError, match="HTTP 402"):  # not a ProviderError: no fallback may hide it
        await provider.structured("decomposition", {"query": "x"}, Understanding)
    memory = Memory(settings(), InMemoryGraph(), provider)
    result = await memory.query("anything")
    assert result["status"] == "failed" and "decomposition failed after 1 attempt(s)" in result["reason"] and "402" in result["reason"]
    await provider.close()


class Unsatisfied(DemoModels):
    """Coverage is never enough, so every query would search online again without search memory."""
    async def structured(self, operation, payload, schema, query_id=None):
        if operation == "coverage":
            return Coverage(coverage=[CoverageItem(information_need_id=n["id"], status="MISSING", evidence_ids=[], missing="more")
                                      for n in payload["information_needs"]], overall_status="INSUFFICIENT")
        return await super().structured(operation, payload, schema, query_id)


async def test_search_memory_skips_needs_searched_recently():
    async def searches(memory, query):
        result = await memory.query(query)
        events = await memory.events(result["query_id"])
        return (sum(e["event_type"] == "EXTERNAL_SEARCH_STARTED" for e in events),
                [e["metadata"] for e in events if e["metadata"].get("reason") == "searched recently"])
    # Episodes would answer the repeats outright; this test is about the online search they would otherwise make.
    memory = Memory(settings(search_memory_similarity=0.85, episode_reuse_hours=0), InMemoryGraph(), Unsatisfied(), DemoExternal())
    await seed(memory)
    assert (await searches(memory, "practice exercises for rootless voicings"))[0] == 1
    started, skipped = await searches(memory, "rootless voicings: practice exercises")  # same content words
    assert started == 0 and skipped[0]["similarity"] == 1.0
    started, skipped = await searches(memory, "practice exercises for rootless voicings today")  # paraphrase, by embedding
    assert started == 0 and 0.85 <= skipped[0]["similarity"] < 1
    assert (await memory.metrics())["searches"] == {"needs": 1, "skipped_as_recent": 2, "reached_context": 1}
    memory.settings.search_memory_days = 0  # disabled (or expired): the web is searched again
    assert (await searches(memory, "practice exercises for rootless voicings"))[0] == 1


async def test_dossier_explores_wider_and_reports_what_the_sources_say():
    from graph_memory.models import QueryRequest
    class Spy(DemoModels):
        operations = []
        async def structured(self, operation, payload, schema, query_id=None):
            self.operations.append(operation)
            return await super().structured(operation, payload, schema, query_id)
    memory = Memory(settings(), InMemoryGraph(), Spy())
    await seed(memory)
    expanded = {}
    for mode in ("answer", "dossier"):
        result = await memory.query(QueryRequest(query="rootless voicings", allow_external=False, mode=mode))
        expanded[mode] = sum(e["event_type"] == "NODE_EXPANDED" for e in await memory.events(result["query_id"]))
    assert result["evidence_ids"] and "dossier" in Spy.operations
    assert expanded["dossier"] > expanded["answer"] == 0  # concepts are followed even once evidence is found


def test_dossier_context_puts_each_sources_best_passage_first():
    from graph_memory.models import Need
    def item(i, document, score):
        return Evidence(id=i, information_need_ids=["N1"], source_node_id=i, source_type="chunk", text=i,
                        relevance_score=score, confidence=1, provenance={"sources": [], "documents": [document]})
    evidence = [item("a1", "A", .9), item("a2", "A", .8), item("b1", "B", .5)]
    builder = ContextBuilder(5000)
    assert [e["id"] for e in builder.build([Need(id="N1", description="x")], evidence)] == ["a1", "a2", "b1"]
    assert [e["id"] for e in builder.build([Need(id="N1", description="x")], evidence, by_source=True)] == ["a1", "b1", "a2"]


async def test_expansion_goes_one_hop_even_without_evidence():
    from graph_memory.models import Decision, Need
    from graph_memory.retrieval import ExplorationSupervisor
    from graph_memory.trace import TraceService
    class ExpandAll(DemoModels):
        async def navigate(self, payload, query_id):
            return NavigationDecision(decisions=[Decision(node_id=c["node_id"], action="EXPAND") for c in payload["candidate_children"]],
                                      current_node_action="CONTINUE")
    repo = InMemoryGraph()  # a chain of concepts, as in a thesaurus: c0 > c1 > c2 > c3
    chain = [Concept(id=f"c{i}", label="Piano" if i == 0 else f"unrelated {i}", preferred_label=str(i)) for i in range(4)]
    await repo.put(chain, [Edge(source=f"c{i}", target=f"c{i+1}", relation=Relation.BROADER_THAN) for i in range(3)])
    for breadth in (False, True):
        trace = TraceService(repo, f"q{breadth}")
        await ExplorationSupervisor(settings(candidate_limit=1), repo, ExpandAll(), trace, breadth=breadth).explore([Need(id="N1", description="piano")])
        events = await repo.records("event", f"q{breadth}:")
        assert [e["node_id"] for e in events if e["event_type"] == "NODE_EXPANDED"] == ["c0"]
        assert any(e["node_id"] == "c1" and e["metadata"].get("reason") == "expansion_depth" for e in events)


async def test_private_scopes_are_isolated_and_stay_out_of_the_concept_graph():
    from graph_memory.models import QueryRequest
    memory = make_memory()
    note = "Piano practice diary: rootless voicings in C, slow tempo, with a metronome."
    private = await memory.ingest(IngestRequest(title="Diary", text=note, scope="user:alice"))
    shared = await memory.ingest(IngestRequest(title="Diary", text=note))
    assert private["document_id"] != shared["document_id"] and not private["duplicate"] and not shared["duplicate"]
    assert not any(e.source == private["document_id"] and e.relation == Relation.ABOUT for e in memory.repository.edges.values())
    async def documents(scope):
        result = await memory.query(QueryRequest(query="piano practice diary rootless voicings", allow_external=False, scope=scope))
        return {d for e in result["evidence"] for d in e["provenance"]["documents"]}
    assert private["document_id"] in await documents("user:alice")
    assert private["document_id"] not in await documents("user:bob") | await documents(None)
    with pytest.raises(ValidationError):
        QueryRequest(query="x", scope="alice")  # a private scope is <kind>:<id>


async def test_forgetting_excludes_restores_erases_and_is_never_refetched():
    from graph_memory.models import QueryRequest
    memory = make_memory()
    await seed(memory)
    question = QueryRequest(query="What is the purpose of rootless voicings?", allow_external=False)
    handbook = next(n.id for n in memory.repository.nodes.values() if n.kind == "Document" and n.label == "Jazz piano handbook")
    async def found():
        return {d for e in (await memory.query(question))["evidence"] for d in e["provenance"]["documents"]}
    assert handbook in await found()
    await memory.forget("outdated", document_id=handbook)
    assert handbook not in await found()
    await memory.restore(handbook)
    assert handbook in await found()
    await memory.forget("wrong", document_id=handbook, hard=True)
    assert not await memory.repository.document_nodes(handbook) and handbook not in await found()
    # A forgotten address is not fetched again by a query's online search.
    await memory.forget("low quality", url="https://example.org/demo/practice")
    events = await memory.events((await memory.query(QUESTIONS[3]))["query_id"])
    assert not any(e["event_type"] == "EXTERNAL_SOURCE_FOUND" for e in events)


async def test_erasing_a_private_scope_keeps_the_shared_memory():
    memory = make_memory()
    await seed(memory)
    shared = len(memory.repository.nodes)
    await memory.ingest(IngestRequest(title="Diary", text="Private practice notes.", scope="user:alice"))
    await memory.repository.record("fact", "user:alice|f1", {"text": "prefers jazz"})
    assert len(memory.repository.nodes) > shared
    await memory.erase_scope("user:alice")
    assert len(memory.repository.nodes) == shared and not await memory.repository.records("fact", "user:alice|")
    with pytest.raises(ValueError):
        await memory.erase_scope("shared")


async def test_old_query_records_and_traces_are_swept():
    memory = make_memory(trace_retention_days=30)
    await seed(memory)
    result = await memory.query("piano keys")
    assert await memory.events(result["query_id"]) and await memory.sweep() == 0
    await memory.repository.record("query", result["query_id"], {**result, "created_at": "2020-01-01T00:00:00+00:00"})
    assert await memory.sweep() == 1
    assert not await memory.events(result["query_id"]) and not await memory.repository.read_record("query", result["query_id"])


async def test_recent_episodes_answer_repeats_until_memory_changes():
    from graph_memory.models import QueryRequest
    class Counting(DemoModels):
        calls = 0
        async def structured(self, operation, payload, schema, query_id=None):
            Counting.calls += 1
            return await super().structured(operation, payload, schema, query_id)
        async def embed(self, text, query_id=None):
            Counting.calls += 1
            return await super().embed(text, query_id)
    memory = Memory(settings(episode_similarity=0.85), InMemoryGraph(), Counting())
    await seed(memory)
    question = QueryRequest(query="What is the purpose of rootless voicings?", allow_external=False, scope="user:alice")
    first = await memory.query(question)
    Counting.calls = 0
    again = await memory.query(question.model_copy(update={"query": "what is the PURPOSE of rootless voicings"}))  # same words
    assert again["reused_from"]["query_id"] == first["query_id"] and again["context"] == first["context"] and Counting.calls == 0
    paraphrase = await memory.query(question.model_copy(update={"query": "What is the purpose of rootless voicings today?"}))
    assert paraphrase["reused_from"]["query_id"] == first["query_id"] and 0.85 <= paraphrase["reused_from"]["similarity"] < 1
    assert "reused_from" not in await memory.query(question.model_copy(update={"reuse": False}))
    assert "reused_from" not in await memory.query(question.model_copy(update={"mode": "dossier"}))  # other options
    await memory.ingest(IngestRequest(title="New", text="Rootless voicings also free the left hand.", scope="user:alice"))
    fresh = await memory.query(question)
    assert "reused_from" not in fresh  # the scope's memory changed since
    [episode] = [e for e in await memory.episodes("user:alice") if e["mode"] == "answer"]
    assert episode["count"] == 5 and episode["query_id"] == fresh["query_id"] and episode["sources"]
    assert (await memory.episodes("user:alice", "purpose of rootless voicings"))[0]["question"]
    assert not await memory.episodes("user:bob")
    # A question that found nothing is searched again rather than answered with the same empty result.
    nothing = question.model_copy(update={"query": "zqxw vlorp brindle"})
    assert not (await memory.query(nothing))["context"] and "reused_from" not in await memory.query(nothing)
    await memory.erase_scope("user:alice")
    assert not await memory.repository.read_record("query", first["query_id"]) and not await memory.episodes("user:alice")


async def test_user_facts_judge_relevance_and_shape_answers_only_in_their_scope():
    from graph_memory.models import QueryRequest
    from graph_memory.prompts import prompt
    class Spy(DemoModels):
        contexts = []
        async def structured(self, operation, payload, schema, query_id=None):
            if operation in ("relevance", "synthesis"):
                self.contexts.append((operation, payload.get("user_context")))
            return await super().structured(operation, payload, schema, query_id)
    memory = Memory(settings(), InMemoryGraph(), Spy())
    await seed(memory)
    question = QueryRequest(query="What is the purpose of rootless voicings?", allow_external=False, scope="user:alice")
    proposed = await memory.add_fact("user:alice", "preference", "Prefers short answers", "inferred", "proposed")
    await memory.query(question)
    assert Spy.contexts and not any(context for _, context in Spy.contexts)  # a proposal applies only once confirmed
    await memory.update_fact("user:alice", proposed["id"], {"status": "active"})
    await memory.add_fact("user:alice", "constraint", "Plays without a bassist")
    Spy.contexts.clear()
    result = await memory.query(question)
    assert "reused_from" not in result  # facts changed, so the earlier answer is not reused
    assert {op for op, context in Spy.contexts if context == ["preference: Prefers short answers", "constraint: Plays without a bassist"]} == {"relevance", "synthesis"}
    Spy.contexts.clear()
    await memory.query(question.model_copy(update={"scope": None}))
    assert not any(context for _, context in Spy.contexts)  # the shared memory knows no user
    assert "user_context" in prompt("relevance", personal=True) and "user_context" not in prompt("relevance")


async def test_inferred_facts_are_proposals_that_cite_their_questions():
    from graph_memory.models import FactProposals, QueryRequest
    class Inferring(DemoModels):
        async def structured(self, operation, payload, schema, query_id=None):
            if operation == "user_facts":
                assert payload["existing_facts"] == ["Avoids shrimp"]
                return FactProposals(facts=[{"kind": "constraint", "text": "Avoids caramel", "evidence": [1, 2]},
                                            {"kind": "objective", "text": "Invented", "evidence": [99]}])
            return await super().structured(operation, payload, schema, query_id)
    memory = Memory(settings(), InMemoryGraph(), Inferring())
    await seed(memory)
    for text in ("courgette feta without caramel", "salad without caramel"):
        await memory.query(QueryRequest(query=text, allow_external=False, scope="user:alice"))
    await memory.add_fact("user:alice", "constraint", "Avoids shrimp")
    [proposal] = await memory.infer_facts("user:alice")  # the one citing no real question is dropped
    assert (proposal["text"], proposal["status"], proposal["origin"]) == ("Avoids caramel", "proposed", "inferred")
    assert set(proposal["evidence"]) == {"courgette feta without caramel", "salad without caramel"}


def test_user_fact_endpoints_need_a_private_scope():
    memory = make_memory()
    with TestClient(create_app(memory)) as client:
        assert client.get("/memory/scopes/shared/facts").status_code == 422
        fact = client.post("/memory/scopes/user:alice/facts", json={"kind": "constraint", "text": "Avoids red meat"}).json()
        assert client.patch(f"/memory/scopes/user:alice/facts/{fact['id']}", json={"text": "Avoids red meat and pork"}).json()["text"] == "Avoids red meat and pork"
        assert [f["text"] for f in client.get("/memory/scopes/user:alice/facts").json()] == ["Avoids red meat and pork"]
        assert client.delete("/memory/scopes/user:alice/facts/unknown").status_code == 404
        assert client.delete(f"/memory/scopes/user:alice/facts/{fact['id']}").status_code == 200
        assert client.delete("/memory/scopes/user:alice").json()["erased"]


def test_images_are_kept_found_again_private_and_forgotten():
    memory = make_memory()
    data = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"pixels").decode()
    with TestClient(create_app(memory)) as client:
        fetched = client.post("/memory/images", json={"data": data, "origin": "fetched", "url": "https://example.org/food.jpg",
                                                      "title": "Omelette"}).json()
        assert fetched["mime_type"] == "image/png"
        # A shared image is found by its address from any scope, so nobody fetches it again.
        [found] = client.get("/memory/images", params={"url": "https://example.org/food.jpg", "scope": "user:bob"}).json()["images"]
        assert client.get(f"/memory/images/{found['id']}", params={"scope": "user:bob"}).json()["data"] == data
        generated = client.post("/memory/images", json={"data": data, "origin": "generated", "prompt": "Annotated violin bow hold",
                                                        "model": "m", "scope": "user:alice"}).json()
        search = {"q": "violin bow hold", "origin": "generated"}
        [candidate] = client.get("/memory/images", params={**search, "scope": "user:alice"}).json()["images"]
        assert candidate["id"] == generated["id"] and candidate["similarity"] > 0 and "data" not in candidate
        assert client.get("/memory/images", params={**search, "scope": "user:bob"}).json()["images"] == []
        assert client.get(f"/memory/images/{generated['id']}", params={"scope": "user:bob"}).status_code == 404
        svg = base64.b64encode(b"<svg onload='alert(1)'/>").decode()
        assert client.post("/memory/images", json={"data": svg, "origin": "fetched", "url": "https://example.org/a.svg"}).status_code == 422
        assert client.post("/memory/images", json={"data": data, "origin": "generated"}).status_code == 422  # no prompt
        assert client.delete("/memory/scopes/user:alice").json()["erased"]
        assert client.get("/memory/images", params={**search, "scope": "user:alice"}).json()["images"] == []
        assert client.delete(f"/memory/images/{fetched['id']}").json() == {"deleted": fetched["id"]}
        assert client.get(f"/memory/images/{fetched['id']}").status_code == 404
        assert client.delete("/memory/images/not-an-id").status_code == 422


async def test_retrieval_know_how_skips_hosts_that_never_help_and_blocked_ones():
    from itertools import count
    from graph_memory.external import RetrievedSource
    serial = count()
    class TwoHosts:
        async def search(self, query, limit):
            n = next(serial)
            return [RetrievedSource(title="Junk", url=f"https://junk.example/{n}", text=f"Buy cheap watches now, offer {n}.", metadata={"retriever": "web"}),
                    RetrievedSource(title="Good", url=f"https://www.good.example/{n}", text=f"Practice exercises: {query}", metadata={"retriever": "wikipedia"})]
    memory = Memory(settings(host_min_trials=5, episode_reuse_hours=0, search_memory_days=0), InMemoryGraph(), Unsatisfied(), TwoHosts())
    async def rejections(question):
        events = await memory.events((await memory.query(question))["query_id"])
        return {e["metadata"]["url"].split("/")[2]: e["metadata"]["reason"] for e in events if e["event_type"] == "EXTERNAL_SOURCE_REJECTED"}
    for word in ("one", "two", "three", "four", "five"):
        assert (await rejections(f"practice exercises {word}")) == {"junk.example": "not relevant"}
    assert (await rejections("practice exercises six")) == {"junk.example": "host never useful (0 of 5 found sources accepted)"}
    listing = {r["name"]: r for r in await memory.retrieval.knowhow.listing()}
    assert listing["good.example"]["accepted"] == 6 and listing["good.example"]["used"] >= 1 and listing["good.example"]["score"] > listing["junk.example"]["score"]
    assert listing["wikipedia"]["kind"] == "retriever" and listing["wikipedia"]["accepted"] == 6
    await memory.retrieval.knowhow.block("good.example", "paywalled")
    assert (await rejections("practice exercises seven"))["www.good.example"] == "host blocked: paywalled"
    await memory.retrieval.knowhow.block("good.example")  # unblocked
    assert "www.good.example" not in await rejections("practice exercises eight")
